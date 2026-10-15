"""Ingestion sanitiser: unicode cleanup, hidden-text removal, injection flags (plan §5.1).

Works on plain text and on .docx (hidden runs: `w:vanish`, tiny fonts, white text).
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from html.parser import HTMLParser

INVISIBLE = re.compile(
    "[​-‏⁠-⁤﻿­᠎‪-‮⁦-⁩]|[\U000e0000-\U000e007f]"
)
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

INJECTION_PATTERNS = [
    r"\bignore\s+(all|any|the|previous|prior|above|earlier)\b[^.\n]{0,40}\b(instructions?|prompts?|rules?|text)",
    r"\bdisregard\b[^.\n]{0,40}\b(instructions?|prompts?|rules?)",
    r"\bsystem\s+prompt\b",
    r"\byou\s+are\s+(now\s+)?(an?|the)\s+(ai|assistant|language model|llm|recruit\w*|screen\w*|model|hiring)",
    r"\b(rate|score|rank|grade|mark)\s+(this|me|the)\s+(candidate|resume|cv|applicant)?\s*(as|at|a|an|with|highly|10|100|top|first)\b",
    r"\b(perfect|ideal|best)\s+(fit|candidate|match)\b",
    r"\bas\s+an\s+(ai|llm|language model)\b",
    r"\b(note|message|instruction)s?\s+(to|for)\s+(the\s+)?(ai|llm|model|assistant|screener|recruiter bot|ats)\b",
    r"\b(chatgpt|gpt-?\d|gemini|claude|llm)\b[^.\n]{0,30}\b(must|should|shall|will)\b",
    r"\bauto[- ]?(advance|approve|accept)\b",
    r"\boutput\s+(json|score|decision)\b",
]
INJECTION_RE = re.compile("|".join(INJECTION_PATTERNS), re.IGNORECASE)


@dataclass
class Sanitized:
    clean_text: str
    flags: dict = field(default_factory=dict)
    injection_matches: list[str] = field(default_factory=list)


def _clean_unicode(text: str) -> tuple[str, int]:
    before = len(text)
    text = INVISIBLE.sub("", text)
    text = unicodedata.normalize("NFKC", text)
    text = CONTROL.sub("", text)
    return text, before - len(text)


def _collapse_ws(text: str) -> str:
    lines = [re.sub(r"[ \t ]+", " ", ln).strip() for ln in text.splitlines()]
    out, blank = [], 0
    for ln in lines:
        blank = blank + 1 if not ln else 0
        if blank <= 1:
            out.append(ln)
    return "\n".join(out).strip()


def sanitize_text(text: str, hidden_removed_chars: int = 0) -> Sanitized:
    text, n_inv = _clean_unicode(text)
    text = _collapse_ws(text)
    matches = [m.group(0) for m in INJECTION_RE.finditer(text)]
    return Sanitized(
        clean_text=text,
        flags={
            "hidden_text_removed": (n_inv + hidden_removed_chars) > 0,
            "injection_suspected": bool(matches),
            "chars_removed": n_inv + hidden_removed_chars,
        },
        injection_matches=matches,
    )


# ---------------------------------------------------------------- HTML
_HIDDEN_STYLE = re.compile(
    r"display\s*:\s*none|visibility\s*:\s*hidden|font-size\s*:\s*([0-3](\.\d+)?)px"
    r"|color\s*:\s*(#fff(fff)?|white|rgb\(\s*255\s*,\s*255\s*,\s*255\s*\))"
    r"|(left|top|text-indent)\s*:\s*-\d{3,}px|opacity\s*:\s*0(\.0+)?(;|$)",
    re.IGNORECASE,
)
_VOID = {"br", "img", "hr", "meta", "link", "input"}


class _VisibleText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.stack: list[bool] = []
        self.parts: list[str] = []
        self.hidden_chars = 0

    def handle_starttag(self, tag, attrs):
        if tag in _VOID:
            if tag == "br":
                self.parts.append("\n")
            return
        a = dict(attrs)
        hidden = bool(_HIDDEN_STYLE.search(a.get("style") or "")) or "hidden" in a or a.get("aria-hidden") == "true"
        self.stack.append(hidden or (self.stack[-1] if self.stack else False) or tag in ("script", "style"))

    def handle_endtag(self, tag):
        if tag in _VOID:
            return
        if self.stack:
            self.stack.pop()
        if tag in ("p", "div", "li", "h1", "h2", "h3", "tr"):
            self.parts.append("\n")

    def handle_data(self, data):
        if self.stack and self.stack[-1]:
            self.hidden_chars += len(data.strip())
        else:
            self.parts.append(data)


def sanitize_html(html: str) -> Sanitized:
    p = _VisibleText()
    p.feed(html)
    return sanitize_text("".join(p.parts), hidden_removed_chars=p.hidden_chars)


# ---------------------------------------------------------------- DOCX
def docx_visible_text(path) -> tuple[str, int]:
    """Return (visible text, hidden chars dropped). Hidden = vanish, <4pt font, white font colour."""
    import docx  # python-docx
    from docx.shared import Pt

    d = docx.Document(str(path))
    hidden = 0
    paras: list[str] = []

    def para_text(p) -> str:
        nonlocal hidden
        out = []
        for r in p.runs:
            f = r.font
            is_hidden = bool(f.hidden) or (f.size is not None and f.size < Pt(4))
            try:
                rgb = f.color.rgb if f.color is not None and f.color.type is not None else None
            except Exception:
                rgb = None
            if rgb is not None and str(rgb).upper() in ("FFFFFF", "FEFEFE", "FDFDFD"):
                is_hidden = True
            if is_hidden:
                hidden += len(r.text.strip())
            else:
                out.append(r.text)
        # paragraphs without runs (rare) fall back to .text
        return "".join(out) if p.runs else p.text

    for p in d.paragraphs:
        paras.append(para_text(p))
    for t in d.tables:
        for row in t.rows:
            paras.append(" | ".join(c.text for c in row.cells))
    return "\n".join(paras), hidden


def sanitize_docx(path) -> Sanitized:
    text, hidden = docx_visible_text(path)
    return sanitize_text(text, hidden_removed_chars=hidden)


# ---------------------------------------------------------------- PDF
def _near_white(color) -> bool:
    """pdfplumber colours: gray (g,), RGB (r,g,b), CMYK (c,m,y,k), or None / pattern names."""
    if color is None or isinstance(color, str):
        return False
    c = tuple(color) if isinstance(color, (list, tuple)) else (color,)
    if not c or not all(isinstance(v, (int, float)) for v in c):
        return False
    if len(c) == 1:
        return c[0] >= 0.95
    if len(c) == 3:
        return min(c) >= 0.95
    if len(c) == 4:
        return max(c) <= 0.05  # CMYK white = no ink
    return False


def pdf_visible_text(path_or_file) -> tuple[str, int, int]:
    """Return (visible text, hidden chars dropped, pages). Hidden = white, <4pt, or placed off the page.

    A PDF with no text layer (a scan) returns empty text: it is escalated as SPARSE_RESUME, never scored from OCR guesses.
    """
    import pdfplumber

    hidden = 0
    pages_text: list[str] = []
    with pdfplumber.open(path_or_file) as pdf:
        n_pages = len(pdf.pages)
        for page in pdf.pages:
            w, h = page.width, page.height

            def off_page(o, w=w, h=h) -> bool:
                return o["x1"] < 0 or o["x0"] > w or o["bottom"] < 0 or o["top"] > h

            # a line that starts on the page and merely overflows the margin is sloppy layout, not hiding:
            # only text whose whole line sits off the page counts as hidden
            on_page_lines = {round(ch["top"]) for ch in page.chars if not off_page(ch)}

            def visible(obj, w=w, h=h) -> bool:
                if obj.get("object_type") != "char":
                    return True
                off = off_page(obj) and round(obj["top"]) not in on_page_lines
                tiny = (obj.get("size") or 12) < 4
                return not (off or tiny or _near_white(obj.get("non_stroking_color")))

            for ch in page.chars:
                if not visible(ch) and ch.get("text", "").strip():
                    hidden += 1
            # word gaps relative to font size: many PDFs contain no space characters at all, and the absolute
            # default (3pt) glues words together ("SoftwareDeveloper", "Dec2025")
            pages_text.append(page.filter(visible).extract_text(x_tolerance_ratio=0.15) or "")
    return "\n".join(t for t in pages_text if t.strip()), hidden, n_pages


def sanitize_pdf(path_or_file) -> Sanitized:
    text, hidden, _ = pdf_visible_text(path_or_file)
    return sanitize_text(text, hidden_removed_chars=hidden)
